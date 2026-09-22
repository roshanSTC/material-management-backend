import io
import re
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Dict, List, Optional, Tuple, Union

import openpyxl

from app.extensions.database import db
from app.models.project import Project


CURRENCY_MAP = {
    "EUR": "€",
    "USD": "$",
    "INR": "₹",
    "GBP": "£",
    "JPY": "¥",
    "AED": "AED",
    "CAD": "CA$",
    "AUD": "AU$",
    "CHF": "CHF",
}

SYMBOL_TO_CURRENCY = {
    "€": "EUR",
    "$": "USD",
    "₹": "INR",
    "£": "GBP",
    "¥": "JPY",
}


def _clean_num(val: Any) -> Optional[Union[int, float]]:
    if val is None or val == "":
        return None
    if isinstance(val, (int, float, Decimal)):
        f = float(val)
        return int(f) if f.is_integer() else round(f, 2)
    s = str(val).replace(",", "").strip()
    try:
        f = float(s)
        return int(f) if f.is_integer() else round(f, 2)
    except ValueError:
        return None


def _clean_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def _parse_date_str(val: Any) -> str:
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    if not val:
        return datetime.utcnow().strftime("%Y-%m-%d")
    s = str(val).strip()
    for fmt in (
        "%d.%m.%Y",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%Y/%m/%d",
        "%m/%d/%Y",
        "%d %b %Y",
        "%d %B %Y",
    ):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return s


def _detect_currency(text: str) -> Tuple[Optional[str], Optional[str]]:
    if not text:
        return None, None
    for code, sym in CURRENCY_MAP.items():
        if f"({code})" in text.upper() or f" {code}" in text.upper() or f"/{code}" in text.upper():
            return code, sym
    for sym, code in SYMBOL_TO_CURRENCY.items():
        if sym in text:
            return code, sym
    return None, None


def _find_target_sheet(wb: openpyxl.Workbook) -> openpyxl.worksheet.worksheet.Worksheet:
    # 1. First priority: Check cell A1/A2/A3 of all sheets for 'Partner Quote'
    for name in wb.sheetnames:
        ws = wb[name]
        for row_idx in (1, 2, 3):
            cell_val = str(ws.cell(row_idx, 1).value or "").strip().lower()
            if "partner quote" in cell_val:
                return ws

    # 2. Priority: Sheet name explicitly containing 'partner quote' or 'partner_quote'
    for name in wb.sheetnames:
        lower = name.lower()
        if "partner quote" in lower or "partner_quote" in lower:
            return wb[name]

    # 3. Priority: Sheet name with 'partner' or 'supplier'
    for name in wb.sheetnames:
        lower = name.lower()
        if "partner" in lower or "supplier" in lower:
            return wb[name]

    # 4. Priority: Sheet name with 'quote' or 'quotation' (avoid customer quotes like 'st-')
    for name in wb.sheetnames:
        lower = name.lower()
        if ("quote" in lower or "quotation" in lower) and not lower.startswith("st-"):
            return wb[name]

    return wb.active or wb.worksheets[0]


def parse_supplier_quotation_excel(
    file_or_stream: Any,
    project_id: Optional[int] = None,
    supplier_id: Optional[int] = None,
) -> Dict[str, Any]:
    """Parse a supplier quotation Excel workbook into frontend form data."""
    if hasattr(file_or_stream, "read"):
        content = file_or_stream.read()
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    elif isinstance(file_or_stream, bytes):
        wb = openpyxl.load_workbook(io.BytesIO(file_or_stream), data_only=True)
    else:
        wb = openpyxl.load_workbook(file_or_stream, data_only=True)

    ws = _find_target_sheet(wb)

    meta: Dict[str, Any] = {}
    table_header_row: Optional[int] = None
    header_cols: Dict[str, int] = {}
    detected_currency: Optional[str] = None
    detected_symbol: Optional[str] = None

    # Scan for metadata and table header
    max_scan_row = min(40, ws.max_row)
    for r in range(1, max_scan_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, min(15, ws.max_column + 1))]
        non_empty = [
            (c, v)
            for c, v in enumerate(row_vals, start=1)
            if v is not None and str(v).strip() != ""
        ]
        if not non_empty:
            continue

        first_col_idx, first_val = non_empty[0]
        first_str = str(first_val).strip()
        first_lower = first_str.lower()

        # Check if this row is the items table header
        row_texts = [str(v).lower() for _, v in non_empty]
        has_desc = any(
            any(k in t for k in ["description", "material name", "item name"])
            for t in row_texts
        )
        has_price_or_qty = any(
            any(k in t for k in ["price", "unit", "qty", "quantity", "rate"])
            for t in row_texts
        )

        if (
            has_desc
            and has_price_or_qty
            and (len(non_empty) >= 3 or "item number" in first_lower)
        ):
            table_header_row = r
            for c in range(1, ws.max_column + 1):
                val = str(ws.cell(r, c).value or "").strip()
                v_low = val.lower()
                if not val:
                    continue

                code, sym = _detect_currency(val)
                if code:
                    detected_currency = code
                    detected_symbol = sym

                if any(k in v_low for k in ["material name", "item description"]) or (
                    "description" in v_low and "desc" not in header_cols
                ):
                    header_cols["desc"] = c
                elif any(
                    k in v_low
                    for k in [
                        "hsn",
                        "hsn_code",
                        "hsn code",
                        "hsn/sac",
                        "sac",
                        "tariff",
                    ]
                ):
                    header_cols["hsn"] = c
                elif any(
                    k in v_low
                    for k in [
                        "material number",
                        "material no",
                        "material_number",
                        "item code",
                        "part number",
                        "part no",
                    ]
                ):
                    header_cols["mat_num"] = c
                elif any(
                    k in v_low
                    for k in [
                        "price per unit",
                        "unit price",
                        "price/unit",
                        "rate",
                    ]
                ):
                    header_cols["price"] = c
                elif any(
                    k in v_low
                    for k in [
                        "net price",
                        "net amount",
                        "total price",
                        "total amount",
                        "amount",
                    ]
                ):
                    header_cols["net"] = c
                elif any(
                    k in v_low
                    for k in ["units", "quantity", "qty", "nos", "count"]
                ):
                    header_cols["units"] = c
            break
        else:
            # Metadata key-value scanning
            if len(non_empty) >= 2:
                key = first_lower
                val = non_empty[1][1]
                meta[key] = val
                code, sym = _detect_currency(str(val))
                if code and not detected_currency:
                    detected_currency = code
                    detected_symbol = sym
            elif len(non_empty) == 1 and ":" in first_str:
                k, v = first_str.split(":", 1)
                meta[k.strip().lower()] = v.strip()

    # Extract items
    items: List[Dict[str, Any]] = []
    total_from_table: Optional[Union[int, float]] = None

    if table_header_row:
        for r in range(table_header_row + 1, ws.max_row + 1):
            row_vals = [
                ws.cell(r, c).value for c in range(1, ws.max_column + 1)
            ]
            if not any(v is not None and str(v).strip() != "" for v in row_vals):
                continue

            # Check if this row is a Total row
            is_total_row = False
            for c, v in enumerate(row_vals, start=1):
                if v and "total" in str(v).lower():
                    is_total_row = True
                    for next_v in row_vals[c:]:
                        amt = _clean_num(next_v)
                        if amt is not None and amt > 0:
                            total_from_table = amt
                            break
                    break

            if is_total_row:
                break

            desc_idx = header_cols.get("desc", 2)
            desc_val = (
                ws.cell(r, desc_idx).value if desc_idx <= ws.max_column else None
            )
            desc_str = str(desc_val or "").strip()
            if not desc_str or "total" in desc_str.lower():
                continue

            mat_num_idx = header_cols.get("mat_num")
            mat_num_val = (
                ws.cell(r, mat_num_idx).value
                if mat_num_idx and mat_num_idx <= ws.max_column
                else ""
            )
            mat_num_str = _clean_str(mat_num_val)

            hsn_idx = header_cols.get("hsn")
            hsn_val = (
                ws.cell(r, hsn_idx).value
                if hsn_idx and hsn_idx <= ws.max_column
                else ""
            )
            hsn_str = _clean_str(hsn_val)

            price_idx = header_cols.get("price")
            price_val = (
                _clean_num(ws.cell(r, price_idx).value)
                if price_idx and price_idx <= ws.max_column
                else None
            )

            units_idx = header_cols.get("units")
            units_val = (
                _clean_num(ws.cell(r, units_idx).value)
                if units_idx and units_idx <= ws.max_column
                else None
            )

            net_idx = header_cols.get("net")
            net_val = (
                _clean_num(ws.cell(r, net_idx).value)
                if net_idx and net_idx <= ws.max_column
                else None
            )

            qty = units_val if units_val is not None else 1
            unit_price = price_val if price_val is not None else 0

            if net_val is None and qty is not None and unit_price is not None:
                net_val = _clean_num(qty * unit_price)

            items.append(
                {
                    "material_name": desc_str,
                    "material_number": mat_num_str,
                    "hsn_code": hsn_str,
                    "quantity": qty,
                    "unit_price": unit_price,
                    "net_amount": net_val if net_val is not None else 0,
                }
            )

    sum_net = sum(it["net_amount"] for it in items)
    final_total_net = (
        total_from_table if total_from_table is not None else _clean_num(sum_net)
    )
    if final_total_net is None:
        final_total_net = 0

    # Currency resolution
    currency_unit = detected_currency or meta.get("currency") or "USD"
    currency_unit = str(currency_unit).strip().upper()
    currency_symbol = (
        detected_symbol or CURRENCY_MAP.get(currency_unit, "$")
    )

    # Quotation Number
    q_num = (
        meta.get("quotation")
        or meta.get("quotation number")
        or meta.get("quotation no")
        or meta.get("quotation_no")
        or meta.get("quote no")
        or meta.get("ref")
        or ""
    )

    # Quotation Date
    q_date = _parse_date_str(
        meta.get("date")
        or meta.get("quotation date")
        or meta.get("quotation_date")
    )

    # Validity
    validity = str(meta.get("validity") or meta.get("price validity") or "").strip()

    # Incoterms
    incoterms = str(
        meta.get("inco terms")
        or meta.get("incoterms")
        or meta.get("inco term")
        or ""
    ).strip()

    # Payment Terms
    payment_terms = str(
        meta.get("payment terms")
        or meta.get("payment term")
        or meta.get("terms of payment")
        or ""
    ).strip()

    # Delivery Period
    delivery_period = str(
        meta.get("delivery time")
        or meta.get("delivery period")
        or meta.get("delivery")
        or ""
    ).strip()

    # Remark (default to empty string)
    remark = str(meta.get("remark") or meta.get("remarks") or "").strip()

    # Quotation Value
    q_val_meta = meta.get("quotation value") or meta.get("total value")
    if q_val_meta is not None:
        clean_q_val = _clean_num(q_val_meta)
        quotation_value_str = (
            str(clean_q_val) if clean_q_val is not None else str(final_total_net)
        )
    else:
        quotation_value_str = str(final_total_net)

    # Resolve project_id and supplier_id
    resolved_project_id = project_id
    if resolved_project_id is None:
        raw_pid = meta.get("project_id") or meta.get("project id")
        if raw_pid is not None:
            try:
                resolved_project_id = int(raw_pid)
            except (ValueError, TypeError):
                resolved_project_id = None

    resolved_supplier_id = supplier_id
    if resolved_supplier_id is None and resolved_project_id:
        try:
            proj = db.session.get(Project, resolved_project_id)
            if proj and proj.supplier_id:
                resolved_supplier_id = proj.supplier_id
        except Exception:
            pass

    if resolved_supplier_id is None:
        raw_sid = meta.get("supplier_id") or meta.get("supplier id")
        if raw_sid is not None:
            try:
                resolved_supplier_id = int(raw_sid)
            except (ValueError, TypeError):
                resolved_supplier_id = None

    return {
        "project_id": resolved_project_id,
        "supplier_id": resolved_supplier_id,
        "quotation_number": str(q_num).strip(),
        "quotation_date": q_date,
        "currency_unit": currency_unit,
        "currency_symbol": currency_symbol,
        "quotation_value": quotation_value_str,
        "total_net_amount": final_total_net,
        "validity": validity,
        "incoterms": incoterms,
        "payment_terms": payment_terms,
        "delivery_period": delivery_period,
        "remark": remark,
        "items": items,
    }
